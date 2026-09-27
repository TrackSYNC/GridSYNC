"""Tests for tools/xlsx_to_schedule.py. Run from the repo root: python3 -m unittest discover tests"""
import datetime as dt
import json
import os
import sys
import tempfile
import unittest

import openpyxl
from openpyxl.styles import PatternFill

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "tools"))
import xlsx_to_schedule as conv  # noqa: E402


def fill(hex_rgb):
    return PatternFill(fill_type="solid", fgColor="FF" + hex_rgb)


def build_sheet(path):
    """A two-day sheet in the same layout as the club's March 2026 schedule."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Friday"
    ws.merge_cells("A1:D1")
    ws["A2"], ws["B2"] = "Time", "Event"
    ws["A3"], ws["B3"] = dt.time(16, 0), "Gates Open"
    ws["A4"], ws["B4"] = "5:30-6:45", "Tech/Registration"
    ws["A5"] = "Saturday"
    ws["A6"], ws["B6"], ws["C6"], ws["D6"] = "Time", "Instructor", "Green", "Blue"
    ws["C6"].fill = fill("C5E0B3")
    ws["A7"], ws["B7"] = "7:45-8:00", "Instructor Meeting - Classroom"
    ws.merge_cells("B7:D7")
    ws["A8"], ws["B8"] = "8:05-8:20", "Driver Meeting - Classroom"
    ws.merge_cells("B8:D8")
    ws["A9"], ws["B9"], ws["D9"] = "8:30-8:55", "Instructor Student Meetup", "Track"
    ws.merge_cells("B9:C9")
    ws["D9"].fill = fill("2E74B5")
    ws["A10"] = "11:45-12:45"
    ws["B10"] = "Lunch / Parade Laps at 12:00pm, (Weather Permitting)"
    ws.merge_cells("B10:D10")
    ws["A11"], ws["C11"] = "1:10-1:35", "Track"
    ws["C11"].fill = fill("538135")
    ws["B11"] = "Check-In"
    ws["A12"] = "not a time"
    wb.save(path)


class ConverterTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.xlsx = os.path.join(self.tmp.name, "sheet.xlsx")
        build_sheet(self.xlsx)

    def tearDown(self):
        self.tmp.cleanup()

    def test_convert(self):
        groups, days, warnings = conv.convert(self.xlsx)
        self.assertEqual([g["id"] for g in groups], ["white", "green", "blue"])
        green = next(g for g in groups if g["id"] == "green")
        self.assertEqual(green["tint"], "#C5E0B3")
        self.assertEqual(green["color"], "#538135")
        self.assertEqual([d["id"] for d in days], ["fri", "sat"])

        fri = days[0]["items"]
        self.assertEqual([(i["start"], i["end"], i["groups"]) for i in fri],
                         [("16:00", "", "all"), ("17:30", "18:45", "all")])

        sat = {i["title"] or (i["kind"] + ":" + ",".join(i["groups"])): i for i in days[1]["items"]}
        self.assertEqual(sat["Instructor Meeting"]["groups"], ["white"])
        self.assertEqual(sat["Instructor Meeting"]["kind"], "meeting")
        self.assertEqual(sat["Instructor Meeting"]["location"], "Classroom")
        self.assertEqual(sat["Driver Meeting"]["groups"], "all")
        self.assertEqual(sat["Instructor Student Meetup"]["groups"], ["white", "green"])
        lunch = sat["Lunch / Parade Laps at 12:00pm"]
        self.assertEqual((lunch["kind"], lunch["start"], lunch["end"], lunch["note"]), ("break", "11:45", "12:45", "Weather Permitting"))
        self.assertEqual(sat["track:green"]["start"], "13:10")
        self.assertEqual(sat["checkin:white"]["end"], "13:35")
        self.assertTrue(any("not a time" in w for w in warnings))

    def test_cli_writes_loadable_file(self):
        out = os.path.join(self.tmp.name, "schedule.js")
        conv.main([self.xlsx, "-o", out, "--name", "Test HPDE", "--dates", "2026-03-06,2026-03-07"])
        with open(out, encoding="utf-8") as fh:
            text = fh.read()
        self.assertIn("window.TSCC_SCHEDULE = {", text)
        data = json.loads(text[text.index("{"):text.rindex("}") + 1])
        self.assertEqual(data["event"]["name"], "Test HPDE")
        self.assertEqual([d["date"] for d in data["days"]], ["2026-03-06", "2026-03-07"])


if __name__ == "__main__":
    unittest.main()
